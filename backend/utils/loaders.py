"""Deterministic extraction with honest source locations.

Character offsets refer to normalized source_text, never original file bytes.
Only PDFs have physical page numbers. OCR and visual layout reconstruction
are outside this extractor's contract.
"""
import hashlib
import io
import json
import re
import codecs
from dataclasses import dataclass, field
from pathlib import Path

EXTRACTOR_VERSION = "structured-text-v1"


class DocumentLoadError(ValueError):
    """A supported file cannot be extracted reliably."""


class UnsupportedFormatError(DocumentLoadError):
    pass


class ScannedPdfError(DocumentLoadError):
    """At least one nonblank PDF page needs OCR before ingestion."""


class EncryptedPdfError(DocumentLoadError):
    pass


class EmptyDocumentError(DocumentLoadError):
    pass


@dataclass(frozen=True)
class SourceBlock:
    text: str
    index: int
    page_number: int | None = None
    section_heading: str | None = None
    start_char: int = 0
    end_char: int = 0


@dataclass
class LoadedDocument:
    doc_type: str
    title: str
    blocks: list[SourceBlock]
    source_hash: str
    source_size: int
    extractor_version: str
    warnings: list[str]
    source_text: str
    page_count: int | None
    _page_texts: list[str] = field(default_factory=list, repr=False)

    @property
    def pages(self) -> list[str]:
        """Compatibility view; non-PDF entries do not imply rendered pages."""
        return list(self._page_texts) if self.doc_type == "pdf" else [self.source_text]

    @classmethod
    def from_text(
        cls, text: str, *, title: str = "untitled", doc_type: str = "txt",
        source_hash: str | None = None, source_size: int | None = None,
    ) -> "LoadedDocument":
        if not isinstance(text, str):
            raise TypeError("text must be a string")
        if doc_type == "pdf":
            raise ValueError("Use from_pages() for PDF text with real page boundaries")
        raw = text.encode("utf-8")
        normalized = clean_text(text)
        headings = _markdown_headings(normalized) if doc_type == "md" else []
        blocks = _text_blocks(normalized, headings=headings)
        return _make_document(
            doc_type, title, blocks, normalized, source_hash or _sha256(raw),
            len(raw) if source_size is None else source_size,
        )

    @classmethod
    def from_pages(
        cls, pages: list[str], *, title: str = "untitled", doc_type: str = "pdf",
        source_hash: str | None = None, source_size: int | None = None,
    ) -> "LoadedDocument":
        if doc_type != "pdf":
            raise ValueError("from_pages() is reserved for real PDF page boundaries")
        if not isinstance(pages, (list, tuple)) or not all(isinstance(p, str) for p in pages):
            raise TypeError("pages must be a sequence of strings")
        # Raw callers supply page boundaries rather than original PDF bytes.
        # An unambiguous representation is essential: joining with a formfeed
        # collides when a page already contains that character. File loaders
        # override this hash/size with those of the original PDF bytes.
        raw = json.dumps(list(pages), ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        normalized_pages = [clean_text(p) for p in pages]
        blocks: list[SourceBlock] = []
        source_parts: list[str] = []
        cursor = 0
        for page_no, text in enumerate(normalized_pages, start=1):
            if source_parts:
                cursor += 2
            source_parts.append(text)
            for block in _text_blocks(text, page_number=page_no):
                blocks.append(SourceBlock(
                    block.text, len(blocks), page_no, block.section_heading,
                    cursor + block.start_char, cursor + block.end_char,
                ))
            cursor += len(text)
        return _make_document(
            doc_type, title, blocks, "\n\n".join(source_parts), source_hash or _sha256(raw),
            len(raw) if source_size is None else source_size,
            page_count=len(pages), page_texts=normalized_pages,
        )


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean_text(text: str) -> str:
    """Keep meaningful layout; normalize line endings, NULs and large gaps."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _make_document(
    doc_type, title, blocks, source_text, source_hash, source_size,
    *, warnings=None, page_count=None, page_texts=None,
) -> LoadedDocument:
    if not source_text.strip() or not blocks:
        raise EmptyDocumentError(f"{title} contains no extractable text")
    return LoadedDocument(
        doc_type=doc_type, title=title, blocks=blocks, source_hash=source_hash,
        source_size=source_size, extractor_version=EXTRACTOR_VERSION,
        warnings=list(warnings or []), source_text=source_text,
        page_count=page_count, _page_texts=list(page_texts or []),
    )


def _markdown_headings(text: str) -> list[tuple[int, str]]:
    """Use the Markdown parser so headings inside fenced code stay code."""
    from markdown_it import MarkdownIt

    lines = text.splitlines(keepends=True)
    starts = [0]
    for line in lines:
        starts.append(starts[-1] + len(line))
    tokens = MarkdownIt("commonmark").parse(text)
    headings = []
    for i, token in enumerate(tokens):
        if token.type == "heading_open" and token.map:
            content = tokens[i + 1].content
            headings.append((starts[token.map[0]], content[:255]))
    return headings


def _text_blocks(
    text: str, *, page_number: int | None = None,
    headings: list[tuple[int, str]] | None = None,
) -> list[SourceBlock]:
    headings = headings or []
    boundaries = {0, len(text)}
    boundaries.update(match.end() for match in re.finditer(r"\n[ \t]*\n", text))
    boundaries.update(position for position, _ in headings)
    positions = sorted(boundaries)
    blocks = []
    for start, end in zip(positions, positions[1:]):
        part = text[start:end]
        left = len(part) - len(part.lstrip())
        right = len(part.rstrip())
        if left >= right:
            continue
        start, end = start + left, start + right
        heading = next((h for pos, h in reversed(headings) if pos <= start), None)
        blocks.append(SourceBlock(text[start:end], len(blocks), page_number, heading, start, end))
    return blocks


def _read_plain(data: bytes, name: str) -> tuple[str, list[str]]:
    for markers, encoding in (
        ((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE), "utf-32"),
        ((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE), "utf-16"),
    ):
        if data.startswith(markers):
            try:
                return data.decode(encoding), [f"Decoded BOM-marked {encoding.upper()} text."]
            except UnicodeDecodeError as exc:
                raise DocumentLoadError(f"{name}: invalid BOM-marked {encoding.upper()} text") from exc
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            warnings = [] if encoding == "utf-8-sig" else ["Decoded using legacy Windows-1252 encoding."]
            return text, warnings
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError(f"{name}: could not decode as UTF-8 or Windows-1252")


def _has_images(page) -> bool:
    """Inspect resource dictionaries without decoding or extracting images."""
    seen: set[int] = set()

    def visit(resources):
        if resources is None:
            return False
        resources = resources.get_object()
        objects = resources.get("/XObject", {})
        objects = objects.get_object() if hasattr(objects, "get_object") else objects
        for ref in objects.values():
            obj = ref.get_object()
            if id(obj) in seen:
                continue
            seen.add(id(obj))
            if obj.get("/Subtype") == "/Image":
                return True
            if obj.get("/Subtype") == "/Form" and visit(obj.get("/Resources")):
                return True
        return False

    return visit(page.get("/Resources"))


def _load_pdf(data: bytes, name: str) -> tuple[list[str], list[str]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise EncryptedPdfError(f"{name}: encrypted PDF requires a password")
        pages, warnings, needs_ocr = [], [], []
        for page_no, page in enumerate(reader.pages, start=1):
            text = clean_text(page.extract_text() or "")
            pages.append(text)
            if not text:
                if _has_images(page):
                    needs_ocr.append(page_no)
                else:
                    warnings.append(f"PDF page {page_no} has no extractable text; treated as blank.")
        if needs_ocr:
            page_list = ", ".join(map(str, needs_ocr))
            raise ScannedPdfError(f"{name}: image-only page(s) {page_list} require OCR; extraction rejected")
        if not any(pages):
            raise EmptyDocumentError(f"{name}: PDF has no extractable text")
        warnings.append("PDF extraction follows the text layer; visual reading order and text inside images are not reconstructed.")
        return pages, warnings
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"{name}: could not read PDF ({type(exc).__name__})") from exc


def _load_docx(data: bytes, title: str, source_hash: str) -> LoadedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
        parts: list[str] = []
        blocks: list[SourceBlock] = []
        cursor = 0
        heading = None
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = clean_text(item.text)
                style = item.style.name if item.style is not None else ""
                if style.lower().startswith("heading") and text:
                    heading = text[:255]
            elif isinstance(item, Table):
                text = clean_text(_table_text(item))
            else:
                continue
            if not text:
                continue
            if parts:
                cursor += 2
            parts.append(text)
            blocks.append(SourceBlock(text, len(blocks), None, heading, cursor, cursor + len(text)))
            cursor += len(text)
        return _make_document(
            "docx", title, blocks, "\n\n".join(parts), source_hash, len(data),
            warnings=["DOCX rendered page numbers, headers, footers, text boxes and tracked revisions are not extracted."],
        )
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"{title}: could not read DOCX ({type(exc).__name__})") from exc


def _table_text(table) -> str:
    # python-docx exposes merged cells at every occupied grid position. Keep
    # their text once, retaining empty placeholders for the occupied columns.
    seen = set()
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            if cell._tc in seen:
                cells.append("")
            else:
                seen.add(cell._tc)
                cells.append(_cell_text(cell))
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def _cell_text(cell) -> str:
    from docx.table import Table

    parts = []
    for item in cell.iter_inner_content():
        if isinstance(item, Table):
            parts.append(_table_text(item))
        else:
            parts.append(item.text)
    return clean_text("\n".join(parts))


def load_document(path) -> LoadedDocument:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if not path.is_file():
        raise IsADirectoryError(f"Not a regular file: {path}")
    doc_type = path.suffix.lower().lstrip(".")
    if doc_type not in {"txt", "md", "pdf", "docx"}:
        raise UnsupportedFormatError(f"Unsupported file type: {path.suffix or '(none)'} ({path.name})")
    data = path.read_bytes()
    source_hash = _sha256(data)
    if doc_type in {"txt", "md"}:
        text, warnings = _read_plain(data, path.name)
        loaded = LoadedDocument.from_text(
            text, title=path.stem, doc_type=doc_type,
            source_hash=source_hash, source_size=len(data),
        )
        loaded.warnings.extend(warnings)
        return loaded
    if doc_type == "pdf":
        pages, warnings = _load_pdf(data, path.name)
        loaded = LoadedDocument.from_pages(
            pages, title=path.stem, source_hash=source_hash, source_size=len(data),
        )
        loaded.warnings.extend(warnings)
        return loaded
    return _load_docx(data, path.stem, source_hash)
